"""Read AlgaTop's product export, so the move to this bot starts where AlgaTop is.

AlgaTop exports one row per product and city with the columns shown in its
cabinet: Артикул, Артикул Kaspi, Название, Статус, Город, Цена, Цена закупа,
Автоснижение, Мин.цена, Автоповышение, Макс.цена, Шаг демпинга, Остатки,
Предзаказ (and a few more that only describe the market, which are ignored).

Columns are found by their header, not by position, and every value is read
strictly: a number column holding text is an error for that row, never a
guess. Both the .xlsx AlgaTop gives and a .csv saved from it are accepted.

Pure module apart from reading the file: no database.
"""

from __future__ import annotations

import csv
import io
import re
from collections.abc import Callable, Iterator
from dataclasses import dataclass
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation

from repricer.cities import CITY_NAMES

MAX_FILE_BYTES = 5_000_000
MAX_ROWS = 20_000

#: Header (lower case, spaces and dots removed) → field.
_COLUMNS = {
    "артикул": "sku",
    "артикулkaspi": "kaspi_product_id",
    "название": "title",
    "статус": "status",
    "город": "city",
    "цена": "price",
    "ценазакупа": "purchase_price",
    "автоснижение": "auto_decrease",
    "минцена": "min_price",
    "автоповышение": "auto_increase",
    "максцена": "max_price",
    "шагдемпинга": "step",
    "остатки": "stock",
    "предзаказ": "preorder",
}
_REQUIRED = ("sku", "city", "price")
_CITY_IDS = {name.lower(): city_id for city_id, name in CITY_NAMES.items()}
_TRUE = {"1", "да", "true", "вкл", "вкл.", "yes"}
_FALSE = {"0", "нет", "false", "выкл", "выкл.", "no", ""}


@dataclass(frozen=True, slots=True)
class AlgaTopRow:
    line: int
    sku: str
    kaspi_product_id: str
    title: str
    #: «Опубликовано» → True, «Снято с продажи» → False, anything else → None.
    published: bool | None
    city_id: str
    price: Decimal
    purchase_price: Decimal | None
    auto_decrease: bool | None
    min_price: Decimal | None
    auto_increase: bool | None
    max_price: Decimal | None
    step: int | None
    stock: int | None
    preorder_days: int | None


@dataclass(frozen=True, slots=True)
class RowError:
    line: int
    sku: str
    reason: str


class AlgaTopFileError(ValueError):
    """The file as a whole cannot be read: wrong format or missing columns."""


def parse_algatop(content: bytes) -> tuple[list[AlgaTopRow], list[RowError]]:
    if len(content) > MAX_FILE_BYTES:
        raise AlgaTopFileError("файл слишком большой (максимум 5 МБ)")
    table = list(_read_table(content))
    header_at = next(
        (index for index, row in enumerate(table[:20]) if "артикул" in {_key(cell) for cell in row}),
        None,
    )
    if header_at is None:
        raise AlgaTopFileError("не нашёл строку заголовков с колонкой «Артикул»")
    columns: dict[str, int] = {}
    for index, cell in enumerate(table[header_at]):
        field = _COLUMNS.get(_key(cell))
        if field is not None and field not in columns:
            columns[field] = index
    missing = [name for name in _REQUIRED if name not in columns]
    if missing:
        names = {"sku": "Артикул", "city": "Город", "price": "Цена"}
        raise AlgaTopFileError("нет колонок: " + ", ".join(names[name] for name in missing))
    body = table[header_at + 1 :]
    if len(body) > MAX_ROWS:
        raise AlgaTopFileError(f"не больше {MAX_ROWS} строк в одном файле")

    rows: list[AlgaTopRow] = []
    errors: list[RowError] = []
    for offset, cells in enumerate(body):
        line = header_at + offset + 2
        if not any(cell.strip() for cell in cells):
            continue

        def value(field: str, cells: list[str] = cells) -> str:
            at = columns.get(field)
            return cells[at].strip() if at is not None and at < len(cells) else ""

        sku = value("sku")
        try:
            rows.append(_row(line, sku, value))
        except ValueError as exc:
            errors.append(RowError(line, sku, str(exc)))
    return rows, errors


def _row(line: int, sku: str, value: Callable[[str], str]) -> AlgaTopRow:
    if not sku or len(sku) > 128:
        raise ValueError("пустой или слишком длинный артикул")
    city = value("city")
    city_id = city if city.isdigit() else _CITY_IDS.get(city.lower())
    if not city_id:
        raise ValueError(f"неизвестный город «{city}»")
    price = _money(value("price"), "цена")
    if price is None:
        raise ValueError("нет цены")
    minimum = _money(value("min_price"), "мин. цена")
    maximum = _money(value("max_price"), "макс. цена")
    if minimum is not None and maximum is not None and minimum > maximum:
        raise ValueError(f"мин. цена {minimum} выше макс. {maximum}")
    kaspi_id = _whole(value("kaspi_product_id"), "артикул Kaspi", 0, None)
    status = value("status").lower()
    return AlgaTopRow(
        line=line,
        sku=sku,
        kaspi_product_id=str(kaspi_id) if kaspi_id else "",
        title=value("title")[:512],
        published=True if status.startswith("опублик") else False if status.startswith("снят") else None,
        city_id=city_id,
        price=price,
        purchase_price=_money(value("purchase_price"), "цена закупа"),
        auto_decrease=_flag(value("auto_decrease"), "автоснижение"),
        min_price=minimum,
        auto_increase=_flag(value("auto_increase"), "автоповышение"),
        max_price=maximum,
        step=_whole(value("step"), "шаг", 0, 1_000_000) or None,
        stock=_whole(value("stock"), "остатки", 0, 2_147_483_647),
        preorder_days=_whole(value("preorder"), "предзаказ", 0, 30),
    )


def _key(cell: str) -> str:
    return re.sub(r"[\s.]+", "", cell.strip().lower())


def _number(raw: str, name: str) -> Decimal | None:
    cleaned = raw.replace(" ", "").replace(" ", "").replace(" ", "").replace("₸", "")
    if not cleaned:
        return None
    try:
        number = Decimal(cleaned.replace(",", "."))
    except InvalidOperation:
        raise ValueError(f"{name}: «{raw}» — не число") from None
    if not number.is_finite() or number < 0:
        raise ValueError(f"{name}: «{raw}» — не число")
    return number


def _money(raw: str, name: str) -> Decimal | None:
    """Whole tenge; 0 and an empty cell both mean "not set" in AlgaTop."""
    number = _number(raw, name)
    if number is None or number == 0:
        return None
    return number.quantize(Decimal(1), rounding=ROUND_HALF_UP)


def _whole(raw: str, name: str, low: int, high: int | None) -> int | None:
    number = _number(raw, name)
    if number is None:
        return None
    if number != number.to_integral_value() or number < low or (high is not None and number > high):
        raise ValueError(f"{name}: «{raw}» — нужно целое число")
    return int(number)


def _flag(raw: str, name: str) -> bool | None:
    text = raw.strip().lower()
    if not text:
        return None
    if text in _TRUE or text in {"1.0"}:
        return True
    if text in _FALSE or text in {"0.0"}:
        return False
    raise ValueError(f"{name}: «{raw}» — ожидается 1 или 0")


def _read_table(content: bytes) -> Iterator[list[str]]:
    if content[:2] == b"PK":
        yield from _read_xlsx(content)
        return
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("cp1251")
    sample = text[:4096]
    delimiter = max(";,\t", key=sample.count)
    yield from csv.reader(io.StringIO(text), delimiter=delimiter)


def _read_xlsx(content: bytes) -> Iterator[list[str]]:
    from openpyxl import load_workbook  # only needed for Excel files

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # zip or XML trouble: not a workbook we can read
        raise AlgaTopFileError("не удалось прочитать Excel-файл") from exc
    try:
        sheet = workbook.worksheets[0]
        for row in sheet.iter_rows(values_only=True):
            yield [_cell(value) for value in row]
    finally:
        workbook.close()


def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)
